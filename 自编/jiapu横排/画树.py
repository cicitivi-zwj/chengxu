# -*- coding: utf-8 -*-
"""
家谱树生成器：读 jiapu.xlsx，画出树形家谱图，存成 家谱树.html。
每次运行都重新读表重新画——Excel 改了，重跑这个脚本（或在家谱程序里按 7），图就是新的。

排列方向（默认横向）：
  横向（默认）——第一代在最左边，第二代往右，一代代往右长；
  竖向（加 --竖）——第一代在最上边，一代代往下长。

支持一个人多次婚姻：表格里「配偶」一列可以填多个编号，用顿号隔开，按结婚先后写，
比如 11、21。同一排框依次排开（本人 → 第一任 → 第二任），
每门婚姻的子女分别挂在自己那门婚姻的后方，各归各家，不会混。
"""

import os
import sys
from datetime import datetime
from openpyxl import load_workbook

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

文件夹 = os.path.dirname(os.path.abspath(__file__))
数据文件 = os.path.join(文件夹, "jiapu.xlsx")
输出文件 = os.path.join(文件夹, "家谱树.html")

# 排列方向：True＝横向（第一代在最左，一代代往右）；False＝竖向（第一代在最上，一代代往下）
# 临时想竖着看，命令行加 --竖：python 画树.py --竖
横排 = "--竖" not in sys.argv

# 也允许在命令行里临时指定：python 画树.py 别的表.xlsx 别的输出.html
# 只给一个文件名时，按扩展名判断：.xlsx 结尾当表格，否则当输出网页
普通参数 = [a for a in sys.argv[1:] if not a.startswith("--")]
if 普通参数:
    if 普通参数[0].lower().endswith((".xlsx", ".xlsm")):
        数据文件 = os.path.abspath(普通参数[0])
        if len(普通参数) > 1:
            输出文件 = os.path.abspath(普通参数[1])
    else:
        输出文件 = os.path.abspath(普通参数[0])

# ── 画布尺寸参数（单位：像素） ──
BOX_W = 110      # 一个人名框的宽
BOX_H = 66       # 高（够放姓名、代数、生卒年三行）
配偶间距 = 30     # 婚姻连线的长度（两个框之间的空隙）
ROW = 120        # 竖向时，上下两代人之间的行距
主轴步 = 150 if 横排 else ROW   # 每往下一代推进多远（横向＝往右，竖向＝往下）
GAP = 28         # 两家人之间的空隙
边距 = 60
顶部 = 40

# 同一个"排"里并排方向上的框尺寸：竖向时排是横着的（看框宽），横向时排是竖着的（看框高）
副轴框 = BOX_H if 横排 else BOX_W

# 颜色：男的框偏蓝，女的框偏粉，没填性别的用灰色
配色 = {
    "男": ("#dbeafe", "#2563eb"),
    "女": ("#fce7f3", "#db2777"),
}


# ── 第一步：读 Excel ──

def 读数据():
    工作簿 = load_workbook(数据文件, data_only=True)
    if "成员" not in 工作簿.sheetnames:
        sys.exit("jiapu.xlsx 里没有名为「成员」的工作表。")
    表 = 工作簿["成员"]
    表头 = [格.value for 格 in 表[1]]
    成员 = []
    for 行 in 表.iter_rows(min_row=2, values_only=True):
        一条 = dict(zip(表头, 行))
        if not 一条.get("姓名") and not 一条.get("编号"):
            continue
        def 规整(x):
            if x is None:
                return ""
            if isinstance(x, datetime):
                return x.strftime("%Y-%m-%d")
            if isinstance(x, float) and x == int(x):
                return str(int(x))
            return str(x).strip()
        成员.append({
            "编号": 规整(一条.get("编号")),
            "姓名": 规整(一条.get("姓名")),
            "性别": 规整(一条.get("性别")),
            "代数": int(float(一条.get("代数") or 0)),
            "父亲": 规整(一条.get("父亲")),
            "母亲": 规整(一条.get("母亲")),
            "配偶": 规整(一条.get("配偶")),
            "阳历出生": 规整(一条.get("阳历出生")),
            "阴历出生": 规整(一条.get("阴历出生")),
            "阳历去世": 规整(一条.get("阳历去世")),
            "阴历去世": 规整(一条.get("阴历去世")),
            "备注": 规整(一条.get("备注")),
        })

    家族名称 = "家谱"
    if "填写说明" in 工作簿.sheetnames:
        for 行 in 工作簿["填写说明"].iter_rows(values_only=True):
            if 行 and 行[0] == "家族名称" and len(行) > 1 and 行[1]:
                家族名称 = str(行[1]).strip()
                break
    return 家族名称, 成员


家族名称, 成员 = 读数据()
按编号 = {p["编号"]: p for p in 成员}


def 取人(编号):
    return 按编号.get(编号)


def 拆分编号(文本):
    """把"11、21"这种写法拆成 ['11', '21']；顿号、逗号、斜杠、空格都认。"""
    if not 文本:
        return []
    文本 = str(文本)
    for 分隔 in "、，,/／;；":
        文本 = 文本.replace(分隔, ",")
    return [x.strip() for x in 文本.replace(" ", ",").split(",") if x.strip()]


def 配偶们(p):
    """这个人的配偶，按先后顺序。表格里自己填的排前面，别人回指的补在后面。"""
    结果 = []
    for 编号 in 拆分编号(p["配偶"]):
        q = 取人(编号)
        if q and q["编号"] not in [x["编号"] for x in 结果]:
            结果.append(q)
    for q in 成员:
        if p["编号"] in 拆分编号(q["配偶"]) and q["编号"] not in [x["编号"] for x in 结果]:
            结果.append(q)
    return 结果


def 子女分组(p):
    """把子女按"和哪一任生的"分堆，返回 (分组列表, 没说清另一位家长的子女)。"""
    配们 = 配偶们(p)
    所有 = [q for q in 成员 if q["父亲"] == p["编号"] or q["母亲"] == p["编号"]]
    分组 = []
    用过 = set()
    for s in 配们:
        这一门 = [q for q in 所有 if q["父亲"] == s["编号"] or q["母亲"] == s["编号"]]
        if 这一门:
            分组.append((s, sorted(这一门, key=lambda x: int(x["编号"]))))
            用过.update(q["编号"] for q in 这一门)
    剩下 = [q for q in 所有 if q["编号"] not in 用过]
    return 分组, sorted(剩下, key=lambda x: int(x["编号"]))


# ── 第二步：搭树。一个节点 = 一个人 + 各任配偶 + 各门子女 ──

已画 = set()

def 建节点(主):
    已画.add(主["编号"])
    配们 = 配偶们(主)
    for s in 配们:
        已画.add(s["编号"])

    合法编号 = {主["编号"]} | {s["编号"] for s in 配们}
    孩子全体 = sorted([q for q in 成员 if q["父亲"] in 合法编号 or q["母亲"] in 合法编号],
                      key=lambda q: int(q["编号"]))

    def 建孩子(k):
        return None if k["编号"] in 已画 else 建节点(k)

    婚 = []
    分掉的 = set()
    for s in 配们:
        这一门 = [k for k in 孩子全体 if k["父亲"] == s["编号"] or k["母亲"] == s["编号"]]
        分掉的.update(k["编号"] for k in 这一门)
        婚.append({"配偶": s, "孩子": [n for n in (建孩子(k) for k in 这一门) if n]})

    剩下 = [k for k in 孩子全体 if k["编号"] not in 分掉的]
    if not 婚:
        # 没记配偶但有子女（单身父亲/母亲），也当成一门
        婚 = [{"配偶": None, "孩子": [n for n in (建孩子(k) for k in 剩下) if n]}]
    else:
        婚[0]["孩子"].extend(n for n in (建孩子(k) for k in 剩下) if n)

    return {"主": 主, "婚": 婚, "w": 0, "主位": 0, "框左": [], "中点": [], "组左": [], "锚点": 0}


根 = []
for p in sorted(成员, key=lambda x: (x["代数"], int(x["编号"]))):
    # 表里没写父母、又还没被画过的，就是一支的带头人
    if not p["父亲"] and not p["母亲"] and p["编号"] not in 已画:
        根.append(建节点(p))

没挂上的人 = [p for p in 成员 if p["编号"] not in 已画]


# ── 第三步：排位置 ──
# 思路：先把本人和几任配偶的框串成一条链，婚姻横线在中点；
#     每门婚姻的子女居中挂在各自中点下方；若和前一组子女挤到一起，就整体往右挪。

def 量宽(node):
    """自下而上：先算孩子需要多宽，再排本节点这一排框和各门子女的位置。"""
    for m in node["婚"]:
        for k in m["孩子"]:
            量宽(k)
        m["孩总"] = sum(k["w"] for k in m["孩子"]) + GAP * (len(m["孩子"]) - 1) if m["孩子"] else 0
    排框(node)


def 排框(node):
    """算出本节点这一排框、各门子女的相对位置（起点从 0 起算），最后整体对齐。"""
    框左 = [0.0]
    中点 = []
    for m in node["婚"]:
        if m["配偶"] is None:
            中点.append(框左[-1] + 副轴框 / 2)          # 单身：孩子挂在本人框正后方
        else:
            中 = 框左[-1] + 副轴框 + 配偶间距 / 2
            中点.append(中)
            框左.append(中 + 配偶间距 / 2)

    组左 = [None] * len(node["婚"])
    前一组右 = None
    for i, m in enumerate(node["婚"]):
        if not m["孩子"]:
            continue
        左 = 中点[i] - m["孩总"] / 2                     # 先试居中
        if 前一组右 is not None and 左 < 前一组右 + GAP:
            左 = 前一组右 + GAP                          # 挤到了就往右挪
        组左[i] = 左
        前一组右 = 左 + m["孩总"]

    最小 = min(框左 + [x for x in 组左 if x is not None])
    框左 = [x - 最小 for x in 框左]
    中点 = [x - 最小 for x in 中点]
    组左 = [None if x is None else x - 最小 for x in 组左]

    右端 = max(框左[-1] + 副轴框,
               max([x + m["孩总"] for x, m in zip(组左, node["婚"]) if x is not None], default=0))

    node["框左"] = 框左
    node["中点"] = 中点
    node["组左"] = 组左
    node["w"] = 右端
    return 右端


行们 = []          # 每个节点实际画在第几行
代数异常 = []      # 代数填得比父母还靠上的人

def 摆放(node, 副轴起点, 主轴起点, 父行=0):
    行 = int(node["主"]["代数"])
    if 行 <= 父行:
        代数异常.append(node["主"])
        行 = 父行 + 1          # 压到下一代去，保证图不叠
    node["行"] = 行
    行们.append(行)
    node["主位"] = 主轴起点 + (行 - 1) * 主轴步
    node["框左"] = [x + 副轴起点 for x in node["框左"]]
    node["中点"] = [x + 副轴起点 for x in node["中点"]]
    node["组左"] = [None if x is None else x + 副轴起点 for x in node["组左"]]
    node["锚点"] = node["框左"][0] + 副轴框 / 2          # 这个节点自己的"落脚点"＝本人框中心
    for i, m in enumerate(node["婚"]):
        if not m["孩子"]:
            continue
        起点 = node["组左"][i]
        for k in m["孩子"]:
            摆放(k, 起点, 主轴起点, 行)
            起点 += k["w"] + GAP


for r in 根:
    量宽(r)
副轴总 = sum(r["w"] for r in 根) + GAP * (len(根) - 1)

起点 = 边距
for r in 根:
    摆放(r, 起点, 顶部)
    起点 += r["w"] + GAP

最大行 = max(行们, default=1)
主轴总 = 顶部 + 最大行 * 主轴步 + 60

# 横向：代际铺在横向上，图就宽；竖向：代际铺在纵向上，图就高
if 横排:
    画布宽, 画布高 = 主轴总, 副轴总 + 边距 * 2
else:
    画布宽, 画布高 = 副轴总 + 边距 * 2, 主轴总


# ── 第四步：画成 SVG ──

def 转义(t):
    return str(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def 屏幕(副轴值, 主轴值):
    """把"第几代的位置、同一排里的位置"换算成屏幕上的 (x, y)。
    竖向：代际往下走，同一排的人左右并排；横向：代际往右走，同一排的人上下并排。"""
    return (主轴值, 副轴值) if 横排 else (副轴值, 主轴值)


def 连线(副1, 主1, 副2, 主2, 色="#9ca3af", 粗=1.5):
    x1, y1 = 屏幕(副1, 主1)
    x2, y2 = 屏幕(副2, 主2)
    return (f'<line x1="{x1:.1f}" y1="{y1:.1f}" x2="{x2:.1f}" y2="{y2:.1f}" '
            f'stroke="{色}" stroke-width="{粗}"/>')


def 生卒简称(p):
    """框里那行小字：1872-1948 / 1949- / 空。"""
    生 = p["阳历出生"][:4] if p["阳历出生"][:4].isdigit() else ""
    去 = p["阳历去世"][:4] if p["阳历去世"][:4].isdigit() else ""
    if 生 and 去:
        return f"{生}-{去}"
    if 生:
        return f"{生}-"
    if 去:
        return f"?-{去}"
    return ""


def 人框(p, 副轴左, 主位):
    """一个人名框。副轴左＝这个人在"排"里的位置，主位＝他属于第几代那一档。"""
    左上x, y = 屏幕(副轴左, 主位)
    填, 边 = 配色.get(p["性别"], ("#e5e7eb", "#6b7280"))
    提示 = [f"{p['姓名']}（第{p['代数']}代）"]
    if p["性别"]:
        提示.append("性别：" + p["性别"])
    if p["阳历出生"] or p["阴历出生"]:
        出生 = []
        if p["阳历出生"]:
            出生.append("阳历 " + p["阳历出生"])
        if p["阴历出生"]:
            出生.append("阴历 " + p["阴历出生"])
        提示.append("出生：" + "　".join(出生))
    if p["阳历去世"] or p["阴历去世"]:
        去世 = []
        if p["阳历去世"]:
            去世.append("阳历 " + p["阳历去世"])
        if p["阴历去世"]:
            去世.append("阴历 " + p["阴历去世"])
        提示.append("去世：" + "　".join(去世))
    if p["父亲"]:
        父 = 取人(p["父亲"])
        提示.append("父亲：" + (父["姓名"] if 父 else "编号" + p["父亲"]))
    if p["母亲"]:
        母 = 取人(p["母亲"])
        提示.append("母亲：" + (母["姓名"] if 母 else "编号" + p["母亲"]))
    配们 = 配偶们(p)
    if len(配们) == 1:
        提示.append("配偶：" + 配们[0]["姓名"])
    elif len(配们) > 1:
        for i, s in enumerate(配们, 1):
            提示.append(f"第{i}任配偶：{s['姓名']}")
    婚组, 剩 = 子女分组(p)
    for s, 孩子们 in 婚组:
        提示.append(f"子女（与{s['姓名']}）：" + "、".join(k["姓名"] for k in 孩子们))
    if 剩:
        提示.append("子女：" + "、".join(k["姓名"] for k in 剩))
    if p["备注"]:
        提示.append("备注：" + p["备注"])

    年岁 = 生卒简称(p)
    第三行 = (f'<text x="{左上x + BOX_W / 2:.1f}" y="{y + 58}" text-anchor="middle" '
              f'font-size="11" fill="#6b7280">{年岁}</text>') if 年岁 else ""

    return (
        f'<g><title>{转义(chr(10).join(提示))}</title>'
        f'<rect x="{左上x:.1f}" y="{y:.1f}" width="{BOX_W}" height="{BOX_H}" rx="8" '
        f'fill="{填}" stroke="{边}" stroke-width="1.5"/>'
        f'<text x="{左上x + BOX_W / 2:.1f}" y="{y + 24}" text-anchor="middle" font-size="15" '
        f'font-weight="bold" fill="#1f2937">{转义(p["姓名"])}</text>'
        f'<text x="{左上x + BOX_W / 2:.1f}" y="{y + 41}" text-anchor="middle" font-size="11" '
        f'fill="#6b7280">第{p["代数"]}代</text>'
        f'{第三行}'
        '</g>'
    )


线条 = []
框 = []

def 画(node):
    主 = node["主"]
    婚 = node["婚"]
    主位 = node["主位"]
    主轴框 = BOX_W if 横排 else BOX_H       # 框沿"代际方向"占的长度

    # 本人 + 各任配偶的框，依次排开；每段婚姻一道连线
    for i, m in enumerate(婚):
        if m["配偶"] is None:
            continue
        # 两端各伸进框里 6 像素，被框盖住，看起来正好是从框边连到框边
        线条.append(连线(node["框左"][i] + 副轴框 - 6, 主位 + 主轴框 / 2,
                         node["框左"][i + 1] + 6, 主位 + 主轴框 / 2, "#6b7280", 2))
    框.append(人框(主, node["框左"][0], 主位))
    for i, m in enumerate(婚):
        if m["配偶"] is not None:
            框.append(人框(m["配偶"], node["框左"][i + 1], 主位))

    # 各门婚姻的子女：从自己那门婚姻的挂点，朝下一代的方向引出去
    落点 = 主位 + 主轴框 + (主轴步 - 主轴框) / 2       # 两代之间的中点
    for i, m in enumerate(婚):
        if not m["孩子"]:
            continue
        挂点 = node["中点"][i]
        # 起点取"框中线"的高度——正好落在婚姻线上，线才不会断一截
        线条.append(连线(挂点, 主位 + 主轴框 / 2, 挂点, 落点))
        孩子锚点 = [k["锚点"] for k in m["孩子"]]
        端点 = 孩子锚点 + [挂点]                               # 带上挂点，保证线不断
        线条.append(连线(min(端点), 落点, max(端点), 落点))      # 横跨所有孩子那一段
        for k in m["孩子"]:
            线条.append(连线(k["锚点"], 落点, k["锚点"], k["主位"]))
            画(k)


for r in 根:
    画(r)

# "第几代"的标尺：竖向时贴在左边，横向时贴在顶上
标尺 = []
for 代 in range(1, 最大行 + 1):
    位置 = 顶部 + (代 - 1) * 主轴步
    if 横排:
        标尺.append(f'<text x="{位置 + BOX_W / 2:.1f}" y="{顶部 - 14:.1f}" text-anchor="middle" '
                    f'font-size="13" fill="#9ca3af">第{代}代</text>')
    else:
        标尺.append(f'<text x="12" y="{位置 + BOX_H / 2:.1f}" font-size="13" '
                    f'fill="#9ca3af">第{代}代</text>')

svg = (f'<svg xmlns="http://www.w3.org/2000/svg" width="{画布宽:.0f}" height="{画布高:.0f}" '
       f'viewBox="0 0 {画布宽:.0f} {画布高:.0f}" font-family="Microsoft YaHei, sans-serif">'
       f'<rect width="100%" height="100%" fill="#fdfdfb"/>' + "".join(标尺) + "".join(线条) + "".join(框) + '</svg>')

警告 = ""
if 没挂上的人:
    名单 = "、".join(f"{p['姓名']}（编号{p['编号']}）" for p in 没挂上的人)
    警告 += ('<div style="background:#fef3c7;border:1px solid #f59e0b;border-radius:8px;'
             'padding:12px 16px;margin:12px 0;font-size:14px;">'
             f'<b>以下 {len(没挂上的人)} 人没画进树里</b>——多半是父母/配偶的编号填错了：'
             f'<br>{名单}<br>改好 Excel 后重跑一次本脚本即可。</div>')
if 代数异常:
    名单 = "、".join(f"{p['姓名']}（表里写第{p['代数']}代）" for p in 代数异常)
    警告 += ('<div style="background:#fee2e2;border:1px solid #ef4444;border-radius:8px;'
             'padding:12px 16px;margin:12px 0;font-size:14px;">'
             f'<b>以下 {len(代数异常)} 人的"代数"填得不合理</b>（不比父母大），'
             f'图上已临时往后压一代：<br>{名单}<br>建议核对 Excel 里的代数。</div>')

方向说明 = "横向排列，第一代在最左、一代代往右" if 横排 else "竖向排列，第一代在最上、一代代往下"

html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<title>{家族名称} · 家谱树</title>
<style>
  body {{ font-family: "Microsoft YaHei", sans-serif; background:#fdfdfb; color:#1f2937; margin:24px; }}
  h1 {{ font-size: 22px; margin: 0 0 4px; }}
  .sub {{ color:#6b7280; font-size: 13px; margin-bottom: 16px; }}
  .board {{ overflow:auto; border:1px solid #e5e7eb; border-radius:10px; background:#fff; }}
  .tip {{ color:#9ca3af; font-size:12px; margin-top:10px; }}
</style>
</head>
<body>
<h1>{家族名称} · 家谱树</h1>
<div class="sub">由 jiapu.xlsx 自动生成，共 {len(成员)} 人，{方向说明}。Excel 改完后重新生成，此图随之更新。</div>
{警告}
<div class="board">{svg}</div>
<div class="tip">鼠标悬停在名字框上可以看到详细信息（含哪几个孩子是哪一任所出）。图比屏幕大时，拖动滚动条就能看全。</div>
</body>
</html>"""

with open(输出文件, "w", encoding="utf-8") as f:
    f.write(html)

print(f"家谱树已生成：{输出文件}")
print(f"共画入 {len(成员) - len(没挂上的人)} 人。{方向说明}。")
多次婚姻的人 = [p for p in 成员 if len(配偶们(p)) > 1]
if 多次婚姻的人:
    print("其中多次婚姻的：" + "、".join(f"{p['姓名']}（{len(配偶们(p))}任）" for p in 多次婚姻的人))
if 没挂上的人:
    print(f"注意：有 {len(没挂上的人)} 人没画进树里（编号可能填错了），打开网页能看到是谁。")
if 代数异常:
    print("注意：有 %d 人的代数填得不合理，图上临时压了一行，网页里有提示。" % len(代数异常))
