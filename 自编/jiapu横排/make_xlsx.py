# -*- coding: utf-8 -*-
"""
一次性工具：把 jiapu.json 里的数据搬进 jiapu.xlsx。
以后数据直接在 Excel 里改，这个脚本用不上了，留着只是备查。
"""

import json
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

文件夹 = os.path.dirname(os.path.abspath(__file__))
json路径 = os.path.join(文件夹, "jiapu.json")
xlsx路径 = os.path.join(文件夹, "jiapu.xlsx")

with open(json路径, "r", encoding="utf-8") as f:
    家谱 = json.load(f)

wb = Workbook()

# ── 第一个表：成员 ──
ws = wb.active
ws.title = "成员"
表头 = ["编号", "姓名", "性别", "代数", "父亲", "母亲", "配偶",
        "阳历出生", "阴历出生", "阳历去世", "阴历去世", "备注"]
ws.append(表头)

# 表头加粗、灰底，好看也好认
粗体 = Font(bold=True)
灰底 = PatternFill("solid", fgColor="DDDDDD")
for 列 in range(1, len(表头) + 1):
    格 = ws.cell(row=1, column=列)
    格.font = 粗体
    格.fill = 灰底
    格.alignment = Alignment(horizontal="center")

for p in 家谱["成员"]:
    ws.append([p["编号"], p["姓名"], p["性别"], p["代数"],
               p["父亲"], p["母亲"], p["配偶"],
               p.get("阳历出生", ""), p.get("阴历出生", ""),
               p.get("阳历去世", ""), p.get("阴历去世", ""),
               p["备注"]])

# 列宽按内容调一下，中文不挤
宽 = {"编号": 8, "姓名": 12, "性别": 8, "代数": 8, "父亲": 8, "母亲": 8, "配偶": 8,
      "阳历出生": 14, "阴历出生": 18, "阳历去世": 14, "阴历去世": 18, "备注": 24}
for i, 名 in enumerate(表头, 1):
    ws.column_dimensions[get_column_letter(i)].width = 宽[名]
ws.freeze_panes = "A2"  # 冻结表头，往下翻时第一行不动

# ── 第二个表：填写说明 ──
说明 = wb.create_sheet("填写说明")
rows = [
    ["填写说明"],
    [""],
    ["家族名称", "示例家族（请把这里改成你家的）"],
    [""],
    ["1. 只在「成员」这个表里填人，一个表里一行就是一个人。"],
    ["2. 编号：随便起，全家人不重复就行（建议 1、2、3……往下排）。"],
    ["3. 父亲、母亲、配偶：填对方的「编号」，不是名字！不认识或没有就留空。"],
    ["4. 代数：第一代填 1，第二代填 2……一共 7 代就填到 7。"],
    ["5. 备注随便写：生卒年、籍贯、埋葬地都行，不写留空。"],
    ["6. 行的顺序无所谓，程序按编号找关系，不按行号。"],
    ["7. 别删表头那一行，别改表头的字。"],
]
for r in rows:
    说明.append(r)
说明.column_dimensions["A"].width = 70
说明.column_dimensions["B"].width = 40
说明["A1"].font = Font(bold=True, size=14)

wb.save(xlsx路径)
print("已生成：" + xlsx路径)
print("共写入 %d 人。" % len(家谱["成员"]))
